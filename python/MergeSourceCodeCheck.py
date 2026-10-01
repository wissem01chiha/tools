from pathlib import Path
import pandas as pd
import tkinter as tk
from tkinter import filedialog, messagebox

from openpyxl import Workbook

SUFFIX = "SrcCheckReport.xlsx"
OUTPUT_NAME = "Merged_SrcCheckReport.xlsx"


def select_folder_with_tk() -> str:
    """Open a Tk dialog to let the user pick a folder and return its path."""
    root = tk.Tk()
    root.withdraw()
    root.attributes("-topmost", True)
    folder = filedialog.askdirectory(title="Select the folder to scan")
    root.destroy()
    return folder


def merge_src_check_reports(folder_path: str,
                            suffix: str = SUFFIX,
                            output_name: str = OUTPUT_NAME) -> Path:
    """
    Scan `folder_path` recursively for .xlsx files whose filename ends with `suffix`.
    Each file is expected to contain exactly one worksheet and identical columns.
    Append all rows into a single output Excel file saved inside `folder_path`.

    Notes on performance:
    - Reading uses pandas/openpyxl (fine).
    - Writing uses openpyxl write_only mode to speed up large exports. [1](https://pytutorial.com/handle-large-excel-files-efficiently-python-openpyxl/)[2](https://woteq.com/how-to-write-to-very-large-excel-files-efficiently-with-openpyxl-write_only-mode/)[3](https://pytutorial.com/optimize-memory-usage-in-python-openpyxl/)
    """
    base = Path(folder_path)

    # Find matching files recursively
    files = sorted([p for p in base.rglob("*.xlsx") if p.name.endswith(suffix)])
    if not files:
        raise FileNotFoundError(f"No .xlsx files found ending with '{suffix}'")

    # Read and collect dataframes
    dfs = []
    errors = []

    for f in files:
        try:
            xls = pd.ExcelFile(f, engine="openpyxl")
            sheet_name = xls.sheet_names[0]  # single sheet expected
            df = pd.read_excel(f, sheet_name=sheet_name, engine="openpyxl")
            dfs.append(df)
        except Exception as e:
            errors.append((str(f), str(e)))

    if not dfs:
        raise RuntimeError("All input files failed to read; nothing to merge.")

    merged = pd.concat(dfs, ignore_index=True)

    # Write output using openpyxl streaming (write_only) for speed
    out_path = base / output_name

    wb = Workbook(write_only=True)
    ws = wb.create_sheet(title="Merged")

    # Write header
    ws.append(list(merged.columns))

    # Write rows (streaming)
    # Using itertuples is usually faster than iterrows
    for row in merged.itertuples(index=False, name=None):
        ws.append(list(row))

    # Optional: write an error sheet (does not alter merged data, no extra columns)
    if errors:
        ws_err = wb.create_sheet(title="Errors")
        ws_err.append(["File", "Error"])
        for f, e in errors:
            ws_err.append([f, e])

    wb.save(out_path)
    return out_path


def main():
    folder = select_folder_with_tk()
    if not folder:
        messagebox.showinfo("Info", "No folder selected. Exiting.")
        return

    try:
        out_path = merge_src_check_reports(folder)
        messagebox.showinfo("Done", f"Report created:\n{out_path}")
    except FileNotFoundError as e:
        messagebox.showwarning("No files found", str(e))
    except Exception as e:
        messagebox.showerror("Error", f"Failed to create the report:\n{e}")


if __name__ == "__main__":
    main()